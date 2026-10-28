"""
build_xlsx.py — 生成 Excel：总览（COUNTIFS 公式）/ 产品清单 / 比较指南 / 组合判断 / 重复风险
用法: python src/build_xlsx.py data/products_classified.json output/VitaminAdviser_Top10.xlsx
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from guide_content import GUIDE, OVERVIEW  # noqa: E402

ORDER = ["magnesium", "vitd", "omega3", "multi", "probiotic", "vitb", "vitc", "calcium", "iron", "glucosamine"]
CN = {k: GUIDE[k]["cn"] for k in ORDER}
VERD = {"good": "合理组合", "maybe": "看情况", "hype": "多为营销"}
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD = PatternFill("solid", fgColor="1F3A33")
WRAP = Alignment(wrap_text=True, vertical="top")

items = json.load(open(sys.argv[1], encoding="utf-8"))
out = sys.argv[2]
wb = Workbook()


def sheet(ws, headers, rows, widths):
    """写入表头、数据，设置字体、列宽、冻结与筛选"""
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        for c in row:
            c.font = F
            c.alignment = WRAP
    for c in ws[1]:
        c.font = FB
        c.fill = HEAD
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# 产品清单
prods = []
for i in items:
    if i["non_supp"]:
        continue
    sec = [CN[t] for t in i["tags"] if t in ORDER and t != i["cat"]]
    if i["cat"] not in ORDER and not sec:
        continue
    prods.append([i["name"], i["brand"], CN.get(i["cat"], "其他"), "、".join(sec), i["variant"] or "—",
                  i["format"], i["group"], i["strength"] or "", i["count"], i["rating"], i["reviews"], i["url"]])
prods.sort(key=lambda r: (ORDER.index(next(k for k in ORDER if CN[k] == r[2])) if r[2] != "其他" else 99, -(r[10] or 0)))
ws2 = wb.active
ws2.title = "产品清单"
sheet(ws2, ["产品名称", "品牌", "主成分", "附加成分（前十内）", "变体", "剂型", "人群", "名称标注剂量", "粒数", "评分", "评论数", "链接"],
      prods, [52, 16, 12, 18, 26, 12, 14, 14, 7, 7, 8, 40])
n = len(prods) + 1

# 总览（公式统计，数据变动可自动更新）
ws1 = wb.create_sheet("总览", 0)
rows = []
for idx, k in enumerate(ORDER, start=2):
    rows.append([CN[k], GUIDE[k]["en"],
                 f"=COUNTIFS('产品清单'!$C$2:$C${n},A{idx})",
                 f"=COUNTIFS('产品清单'!$D$2:$D${n},\"*\"&A{idx}&\"*\")",
                 f"=C{idx}+D{idx}", GUIDE[k]["tagline"], GUIDE[k]["test"]])
sheet(ws1, ["成分", "英文", "主成分产品数", "作为附加成分", "合计", "一句话定位", "是否先检查"], rows,
      [18, 22, 13, 13, 9, 60, 60])
ws1.cell(row=len(ORDER) + 3, column=1,
         value="说明：数量由「产品清单」用 COUNTIFS 计算；分类仅依据产品名称。数据来源：chemistwarehouse.csv（用户提供）。").font = F

# 比较指南（长表）
LAB = [("who", "谁可能需要"), ("when", "什么情况下考虑"), ("caution", "谁要谨慎/药物相互作用"),
       ("compare", "比较清单"), ("food", "食物与生活方式替代"), ("myths", "常见误区")]
rows = []
for k in ORDER:
    g = GUIDE[k]
    rows.append([g["cn"], "是否先检查", g["test"], ""])
    for key, lab in LAB:
        for x in g[key]:
            rows.append([g["cn"], lab, x, ""])
    for f in g["forms"]:
        rows.append([g["cn"], "成分形式", f"{f[0]}：{f[1]}", f[2]])
sheet(wb.create_sheet("比较指南"), ["成分", "项目", "内容", "较适合"], rows, [18, 22, 90, 26])

# 组合判断
rows = [[GUIDE[k]["cn"], c[0], VERD[c[1]], c[2]] for k in ORDER for c in GUIDE[k]["combos"]]
sheet(wb.create_sheet("组合判断"), ["成分", "组合", "判断", "说明"], rows, [18, 34, 12, 80])

# 重复风险
sheet(wb.create_sheet("重复风险"), ["成分", "常藏在这些产品里", "叠加的风险"], OVERVIEW["overlap"], [16, 60, 40])
wb.save(out)
print("saved", out, "products", len(prods))
